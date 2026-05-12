from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db.models import Count, Q
from django.db.models.functions import TruncMonth

from .models import User
from placements.models import PlacementDrive, Application, Interview
from students.models import StudentProfile, RecruiterProfile
from notifications.models import Notification
from notifications.views import send_notification
from campus_connect.email_utils import send_status_update_email


def landing_view(request):
    return render(request, 'landing.html')


def register_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        username   = request.POST.get('username', '')
        email      = request.POST.get('email', '')
        password1  = request.POST.get('password1', '')
        password2  = request.POST.get('password2', '')
        first_name = request.POST.get('first_name', '')
        last_name  = request.POST.get('last_name', '')
        phone      = request.POST.get('phone', '')

        role = request.POST.get('role', 'student')
        if role not in ['student', 'recruiter']:
            role = 'student'

        if password1 != password2:
            messages.error(request, 'Passwords do not match!')
            return redirect('register')

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username already taken!')
            return redirect('register')

        user = User.objects.create_user(
            username=username,
            email=email,
            password=password1,
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            role=role
        )

        messages.success(request, 'Account created! Please login.')
        return redirect('login')

    return render(request, 'accounts/register.html')


def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        username = request.POST['username']
        password = request.POST['password']
        user = authenticate(request, username=username, password=password)

        if user is not None:
            if not user.is_active:
                messages.error(request, 'Your account has been deactivated by the administrator.')
                return redirect('login')
            login(request, user)
            User.objects.filter(pk=request.user.pk).update(last_active_at=timezone.now())          
            messages.success(request, f'Welcome back, {user.first_name}!')
            return redirect('dashboard')
        else:
            messages.error(request, 'Invalid username or password!')

    return render(request, 'accounts/login.html')


def logout_view(request):
    logout(request)
    messages.success(request, 'Logged out successfully!')
    return redirect('landing')


@login_required
def dashboard_view(request):
    user = request.user
    context = {'user': user}

    # ── SAFE HEARTBEAT UPDATE ──
    # Only update last_active_at, bypassing any signals that might trigger save() loops
    User.objects.filter(pk=user.pk).update(last_active_at=timezone.now())
    # Refresh the user object so the template has the new data if needed
    user.refresh_from_db()

    open_drives = PlacementDrive.objects.filter(status='active')   
    context['open_drives_count'] = open_drives.count()
    context['recent_drives'] = open_drives.order_by('-created_at')[:5]

    context['notifications'] = Notification.objects.filter(user=user).order_by('-created_at')[:4]
    context['unread_count'] = Notification.objects.filter(user=user, is_read=False).count()

    if user.role == 'student':
        apps = Application.objects.filter(student=user)
        context['applications_count'] = apps.count()
        context['shortlisted_count'] = apps.filter(status='shortlisted').count()
        context['my_applications'] = apps.order_by('-applied_at')[:4]
        context['placed_applications'] = apps.filter(status='placed')

        interviews = Interview.objects.filter(
            application__student=user,
            scheduled_at__gte=timezone.now()
        ).order_by('scheduled_at')
        context['interviews_count'] = interviews.count()
        context['upcoming_interviews'] = interviews[:2]

        try:
            context['profile_completion'] = user.studentprofile.profile_completion
        except StudentProfile.DoesNotExist:
            context['profile_completion'] = 0

    elif user.role == 'recruiter':
        my_drives = PlacementDrive.objects.filter(posted_by=user)
        apps = Application.objects.filter(drive__in=my_drives)
        context['applications_count'] = apps.count()
        context['shortlisted_count'] = apps.filter(status='shortlisted').count()
        context['interviews_count'] = apps.filter(status='shortlisted').count()
        context['my_drives'] = my_drives.order_by('-created_at')[:5]

    elif user.role == 'admin':
        apps = Application.objects.all()
        context['applications_count'] = apps.count()
        context['shortlisted_count'] = apps.filter(status='shortlisted').count()
        context['interviews_count'] = apps.filter(status='shortlisted').count()
        context['total_students'] = User.objects.filter(role='student').count()
        context['total_recruiters'] = User.objects.filter(role='recruiter').count()

    return render(request, 'dashboard.html', context)


@login_required
def admin_dashboard_view(request):
    if request.user.role != 'admin':
        messages.error(request, 'Access denied!')
        return redirect('dashboard')

    students   = User.objects.filter(role='student')
    recruiters = User.objects.filter(role='recruiter').select_related('recruiterprofile')
    drives     = PlacementDrive.objects.all().order_by('-created_at')
    apps       = Application.objects.all().order_by('-applied_at')

    total_students    = students.count()
    total_recruiters  = recruiters.count()
    total_drives      = drives.count()
    total_apps        = apps.count()
    selected_students = apps.filter(status='selected').count()

    pending_recruiters = User.objects.filter(
        role='recruiter',
        recruiterprofile__verification_status='pending'
    ).select_related('recruiterprofile')

    return render(request, 'accounts/admin_dashboard.html', {
        'students':           students,
        'recruiters':         recruiters,
        'drives':             drives,
        'apps':               apps,
        'total_students':     total_students,
        'total_recruiters':   total_recruiters,
        'total_drives':       total_drives,
        'total_apps':         total_apps,
        'selected_students':  selected_students,
        'pending_recruiters': pending_recruiters,
    })


@login_required
def update_application_status(request, pk):
    app = get_object_or_404(Application, pk=pk)

    if request.user.role != 'admin' and app.drive.posted_by != request.user:
        messages.error(request, 'Access denied!')
        return redirect('dashboard')

    if request.method == 'POST':
        status = request.POST.get('status')

        # ✨ 'placed' is now included in this single check
        if status in ['applied', 'shortlisted', 'rejected', 'selected', 'placed']:
            app.status = status
            app.save()

            # ── AUTO-CREATE CONVERSATION ON SHORTLIST ──
            if status == 'shortlisted':
                from placements.models import Conversation # Import here to avoid circular imports
                
                # Get or create prevents duplicates if they un-shortlist and re-shortlist
                conv, created = Conversation.objects.get_or_create(
                    application=app,
                    defaults={
                        'student': app.student,
                        'recruiter': app.drive.posted_by,
                        'drive': app.drive
                    }
                )
                
                # Add a system message to start the chat
                if created:
                    from placements.models import ChatMessage
                    ChatMessage.objects.create(
                        conversation=conv,
                        sender='system',
                        message_type='system',
                        content=f"Chat initiated. {app.drive.posted_by.get_full_name()} has shortlisted you for {app.drive.role}."
                    )

            send_notification(
                user=app.student,
                title=f'Application Update - {app.drive.company}',
                message=f'Your application for {app.drive.role} at {app.drive.company} has been {status}!',
                notif_type=status if status in ['shortlisted', 'rejected', 'selected', 'placed'] else 'general'
            )

            send_status_update_email(app.student, app.drive, status)
            messages.success(request, f'Application status updated to {status}!')

    if request.user.role == 'recruiter':
        return redirect('recruiter_applicants')
    else:
        return redirect('admin_dashboard')        


@login_required
def stats_view(request):
    import json

    user = request.user
    context = {}

    if user.role == 'student':
        my_apps = Application.objects.filter(student=user)
        context['total_applied'] = my_apps.count()
        context['shortlisted'] = my_apps.filter(status='shortlisted').count()
        context['selected'] = my_apps.filter(status='selected').count()
        context['rejected'] = my_apps.filter(status='rejected').count()

        context['status_labels'] = json.dumps(['Applied', 'Shortlisted', 'Selected', 'Rejected'])
        context['status_data'] = json.dumps([
            my_apps.filter(status='applied').count(),
            context['shortlisted'],
            context['selected'],
            context['rejected']
        ])

    elif user.role == 'recruiter':
        my_drives = PlacementDrive.objects.filter(posted_by=user)
        my_apps = Application.objects.filter(drive__in=my_drives)

        context['total_drives'] = my_drives.count()
        context['active_drives'] = my_drives.filter(status='active').count()
        context['total_applicants'] = my_apps.count()
        context['shortlisted'] = my_apps.filter(status='shortlisted').count()
        context['selected'] = my_apps.filter(status='selected').count()

        context['status_labels'] = json.dumps(['Applied', 'Shortlisted', 'Selected', 'Rejected'])
        context['status_data'] = json.dumps([
            my_apps.filter(status='applied').count(),
            context['shortlisted'],
            context['selected'],
            my_apps.filter(status='rejected').count()
        ])

        drive_names = list(my_drives.values_list('role', flat=True))
        drive_applicants = list(my_drives.annotate(app_count=Count('application')).values_list('app_count', flat=True))
        context['drive_labels'] = json.dumps(drive_names)
        context['drive_data'] = json.dumps(drive_applicants)

    elif user.role == 'admin':
        total_students = User.objects.filter(role='student').count()
        total_apps = Application.objects.all().count()
        selected = Application.objects.filter(status='selected').count()

        context['total_students'] = total_students
        context['total_recruiters'] = User.objects.filter(role='recruiter').count()
        context['total_drives'] = PlacementDrive.objects.count()
        context['total_applications'] = total_apps
        context['selected_students'] = selected
        context['placement_rate'] = round((selected / total_students * 100), 1) if total_students else 0

        context['status_labels'] = json.dumps(['Applied', 'Shortlisted', 'Selected', 'Rejected'])
        context['status_data'] = json.dumps([
            Application.objects.filter(status='applied').count(),
            Application.objects.filter(status='shortlisted').count(),
            Application.objects.filter(status='selected').count(),
            Application.objects.filter(status='rejected').count()
        ])

        monthly_data = Application.objects.annotate(month=TruncMonth('applied_at')).values('month').annotate(count=Count('id')).order_by('month')
        context['monthly_labels'] = json.dumps([m['month'].strftime('%b %Y') for m in monthly_data if m['month']])
        context['monthly_data'] = json.dumps([m['count'] for m in monthly_data])

        branch_data = Application.objects.filter(status='selected').values('student__studentprofile__branch').annotate(count=Count('id')).order_by('-count')
        context['branch_labels'] = json.dumps([b['student__studentprofile__branch'] or 'N/A' for b in branch_data])
        context['branch_data'] = json.dumps([b['count'] for b in branch_data])

    return render(request, 'accounts/stats.html', context)


@login_required
def export_students_csv(request):
    if request.user.role != 'admin':
        return redirect('dashboard')

    import csv
    from django.http import HttpResponse

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="students.csv"'
    writer = csv.writer(response)
    writer.writerow(['Name', 'Email', 'Username', 'Branch', 'CGPA',
                     'Roll Number', 'Passing Year', 'Skills', 'LinkedIn', 'GitHub'])

    for user in User.objects.filter(role='student').select_related('studentprofile'):
        try:
            p = user.studentprofile
            writer.writerow([
                user.get_full_name(), user.email, user.username,
                p.branch, p.cgpa, p.roll_number, p.passing_year,
                p.skills, p.linkedin, p.github
            ])
        except StudentProfile.DoesNotExist:
            writer.writerow([user.get_full_name(), user.email, user.username,
                             '', '', '', '', '', '', ''])
    return response


@login_required
def export_students_excel(request):
    if request.user.role != 'admin':
        return redirect('dashboard')

    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = ['Name', 'Email', 'Username', 'Phone', 'Branch', 'CGPA',
               'Roll Number', 'Passing Year', 'Skills', 'LinkedIn', 'GitHub']

    header_fill = PatternFill(start_color="5B6CFF", end_color="5B6CFF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for user in User.objects.filter(role='student').select_related('studentprofile'):
        try:
            p = user.studentprofile
            ws.append([
                user.get_full_name(), user.email, user.username, user.phone,
                p.branch, p.cgpa, p.roll_number, p.passing_year,
                p.skills, p.linkedin, p.github
            ])
        except StudentProfile.DoesNotExist:
            ws.append([user.get_full_name(), user.email, user.username,
                       user.phone, '', '', '', '', '', '', ''])

    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="students.xlsx"'
    wb.save(response)
    return response


@login_required
def search_students_view(request):
    if request.user.role not in ['recruiter', 'admin']:
        return redirect('dashboard')

    students = StudentProfile.objects.select_related('user').all()
    branch   = request.GET.get('branch', '')
    skill    = request.GET.get('skill', '')
    min_cgpa = request.GET.get('min_cgpa', '')

    if branch:
        students = students.filter(branch=branch)
    if skill:
        students = students.filter(skills__icontains=skill)
    if min_cgpa:
        try:
            students = students.filter(cgpa__gte=float(min_cgpa))
        except ValueError:
            pass

    return render(request, 'accounts/search_students.html', {
        'students': students,
        'branch':   branch,
        'skill':    skill,
        'min_cgpa': min_cgpa,
    })


# ── VERIFICATION VIEWS ────────────────────────────────────────────────────

@login_required
def verification_status_view(request):
    if request.user.role != 'recruiter':
        return redirect('dashboard')
    profile, _ = RecruiterProfile.objects.get_or_create(user=request.user)
    return render(request, 'accounts/verification_status.html', {'profile': profile})


@login_required
def request_verification_view(request):
    if request.user.role != 'recruiter':
        return redirect('dashboard')

    profile, _ = RecruiterProfile.objects.get_or_create(user=request.user)

    if profile.is_verified:
        return redirect('dashboard')

    if request.method == 'POST':
        profile.company_name              = request.POST.get('company_name', '')
        profile.designation               = request.POST.get('designation', '')
        profile.industry                  = request.POST.get('industry', '')
        profile.company_website           = request.POST.get('company_website', '')
        profile.verification_note         = request.POST.get('verification_note', '')
        profile.verification_status       = 'pending'
        profile.verification_requested_at = timezone.now()
        profile.save()

        messages.success(request, 'Verification request submitted! We will review it shortly.')
        return redirect('verification_status')

    return render(request, 'accounts/request_verification.html', {'profile': profile})


@login_required
def admin_verify_recruiter(request, user_id):
    if request.user.role != 'admin':
        messages.error(request, 'Access denied.')
        return redirect('dashboard')

    recruiter_user = get_object_or_404(User, pk=user_id, role='recruiter')
    profile, _     = RecruiterProfile.objects.get_or_create(user=recruiter_user)

    if request.method == 'POST':
        action       = request.POST.get('action')
        admin_remark = request.POST.get('admin_remark', '')
        profile.admin_remark = admin_remark

        if action == 'approve':
            profile.is_verified         = True
            profile.verification_status = 'approved'
            profile.verified_at         = timezone.now()
            profile.save()

            recruiter_user.is_verified = True
            recruiter_user.save()

            send_notification(
                user=recruiter_user,
                title='Verification Approved!',
                message='Your recruiter account has been verified. You can now post placement drives.',
                notif_type='general'
            )
            messages.success(request, f'{recruiter_user.username} has been verified.')

        elif action == 'reject':
            profile.is_verified         = False
            profile.verification_status = 'rejected'
            profile.save()

            recruiter_user.is_verified = False
            recruiter_user.save()

            send_notification(
                user=recruiter_user,
                title='Verification Update',
                message=f'Your verification request was not approved. Reason: {admin_remark}',
                notif_type='general'
            )
            messages.warning(request, f'{recruiter_user.username} verification rejected.')

    return redirect('admin_dashboard')


# ── ADMIN DRIVE MODERATION ────────────────────────────────────────────────

@login_required
def admin_toggle_drive_status(request, drive_id):
    if request.user.role != 'admin':
        messages.error(request, 'Access denied.')
        return redirect('dashboard')
    
    drive = get_object_or_404(PlacementDrive, pk=drive_id)
    
    if request.method == 'POST':
        if drive.status == 'active':
            drive.status = 'closed'
            messages.warning(request, f'Drive "{drive.company} - {drive.role}" has been forcefully closed.')
        else:
            drive.status = 'active'
            messages.success(request, f'Drive "{drive.company} - {drive.role}" has been reopened.')
        
        drive.save()
        
    return redirect('admin_dashboard')


# ── ADMIN USER MANAGEMENT (SOFT DELETE) ───────────────────────────────────

@login_required
def admin_toggle_user_status(request, user_id):
    if request.user.role != 'admin':
        messages.error(request, 'Access denied.')
        return redirect('dashboard')
    
    target_user = get_object_or_404(User, pk=user_id)
    
    # Prevent admin from deactivating themselves
    if target_user == request.user:
        messages.error(request, "You cannot deactivate your own account!")
        return redirect('admin_dashboard')
        
    if request.method == 'POST':
        # Toggle the is_active status
        target_user.is_active = not target_user.is_active
        target_user.save()
        
        if target_user.is_active:
            messages.success(request, f'User {target_user.username} has been reactivated.')
        else:
            messages.warning(request, f'User {target_user.username} has been deactivated. They can no longer log in.')
            
    return redirect('admin_dashboard')